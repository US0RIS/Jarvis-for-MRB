from __future__ import annotations

import base64
import io
import json
import re
import sqlite3
import zipfile
from datetime import datetime
from email.utils import parseaddr
from pathlib import PurePath
from typing import Any, Iterable
from xml.etree import ElementTree

from googleapiclient.discovery import build

from jarvis_mrb.world_model import DB_PATH, ensure_entity, record_knowledge_source

_MAX_ATTACHMENT_BYTES = 8 * 1024 * 1024
_MAX_TEXT_CHARS = 80_000
_MAX_MESSAGES = 20
_MAX_PAGES_PER_SYNC = 5
_OVERLAP_SECONDS = 48 * 60 * 60
_ALLOWED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".txt", ".md", ".csv", ".json", ".rtf", ".pptx"}
_DEFAULT_QUERY = "in:anywhere newer_than:30d has:attachment"


def _state_connect() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH, timeout=10.0)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA busy_timeout=10000")
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS gmail_attachment_sync_state (
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL
        )
        """
    )
    conn.commit()
    return conn


def _state_get(key: str, default: str = "") -> str:
    try:
        with _state_connect() as conn:
            row = conn.execute("SELECT value FROM gmail_attachment_sync_state WHERE key=?", (key,)).fetchone()
        return str(row["value"]) if row else default
    except sqlite3.Error:
        return default


def _state_set(**values: Any) -> None:
    try:
        with _state_connect() as conn:
            for key, value in values.items():
                conn.execute(
                    "INSERT INTO gmail_attachment_sync_state(key,value) VALUES(?,?) "
                    "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                    (str(key), json.dumps(value, ensure_ascii=False) if not isinstance(value, str) else value),
                )
            conn.commit()
    except sqlite3.Error:
        pass


def _state_int(key: str) -> int:
    raw = _state_get(key, "0")
    try:
        return int(json.loads(raw)) if raw[:1] in "-0123456789" else int(raw)
    except (TypeError, ValueError, json.JSONDecodeError):
        return 0


def _decode_base64url(data: str) -> bytes:
    padded = str(data or "") + "=" * (-len(str(data or "")) % 4)
    return base64.urlsafe_b64decode(padded.encode("ascii"))


def _iter_parts(part: dict[str, Any]) -> Iterable[dict[str, Any]]:
    yield part
    for child in part.get("parts") or []:
        if isinstance(child, dict):
            yield from _iter_parts(child)


def _headers(payload: dict[str, Any]) -> dict[str, str]:
    result: dict[str, str] = {}
    for item in payload.get("headers") or []:
        if not isinstance(item, dict):
            continue
        name = str(item.get("name") or "").strip().lower()
        value = str(item.get("value") or "").strip()
        if name and value:
            result[name] = value
    return result


def _clean_text(value: str) -> str:
    text = value.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:_MAX_TEXT_CHARS]


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "utf-16", "windows-1252", "latin-1"):
        try:
            return _clean_text(data.decode(encoding))
        except UnicodeDecodeError:
            continue
    return ""


def _extract_docx(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = set(archive.namelist())
            targets = ["word/document.xml"]
            targets.extend(sorted(name for name in names if name.startswith("word/header") and name.endswith(".xml")))
            targets.extend(sorted(name for name in names if name.startswith("word/footer") and name.endswith(".xml")))
            chunks: list[str] = []
            for name in targets:
                if name not in names:
                    continue
                raw = archive.read(name)
                root = ElementTree.fromstring(raw)
                text = " ".join(node.text or "" for node in root.iter() if node.tag.endswith("}t"))
                if text.strip():
                    chunks.append(text.strip())
                if sum(len(chunk) for chunk in chunks) >= _MAX_TEXT_CHARS:
                    break
            return _clean_text("\n".join(chunks))
    except (zipfile.BadZipFile, KeyError, ElementTree.ParseError, OSError):
        return ""


def _extract_pptx(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            slide_names = sorted(
                name for name in archive.namelist()
                if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
            )[:120]
            chunks: list[str] = []
            for name in slide_names:
                root = ElementTree.fromstring(archive.read(name))
                text = " ".join(node.text or "" for node in root.iter() if node.tag.endswith("}t"))
                if text.strip():
                    chunks.append(text.strip())
                if sum(len(chunk) for chunk in chunks) >= _MAX_TEXT_CHARS:
                    break
            return _clean_text("\n".join(chunks))
    except (zipfile.BadZipFile, KeyError, ElementTree.ParseError, OSError):
        return ""


def _extract_xlsx(data: bytes) -> str:
    """Passively read cell values; formulas are never executed."""
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        chunks: list[str] = []
        total = 0
        try:
            for sheet in workbook.worksheets[:40]:
                chunks.append(f"Sheet: {sheet.title}")
                total += len(chunks[-1])
                for row in sheet.iter_rows(values_only=True):
                    values = [" ".join(str(value).split()) for value in row if value is not None and str(value).strip()]
                    if not values:
                        continue
                    line = " | ".join(values)[:4000]
                    chunks.append(line)
                    total += len(line)
                    if total >= _MAX_TEXT_CHARS:
                        break
                if total >= _MAX_TEXT_CHARS:
                    break
        finally:
            workbook.close()
        return _clean_text("\n".join(chunks))
    except Exception:
        return ""


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data), strict=False)
        chunks: list[str] = []
        for page in list(reader.pages)[:80]:
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""
            if text.strip():
                chunks.append(text.strip())
            if sum(len(chunk) for chunk in chunks) >= _MAX_TEXT_CHARS:
                break
        return _clean_text("\n".join(chunks))
    except Exception:
        return ""


def _extract_rtf(data: bytes) -> str:
    text = _decode_text(data)
    if not text:
        return ""
    text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", text)
    text = text.replace("{", " ").replace("}", " ")
    return _clean_text(text)


def _extract(filename: str, mime_type: str, data: bytes) -> str:
    extension = PurePath(filename).suffix.lower()
    if extension not in _ALLOWED_EXTENSIONS:
        return ""
    if len(data) > _MAX_ATTACHMENT_BYTES:
        return ""
    if extension == ".pdf":
        return _extract_pdf(data)
    if extension == ".docx":
        return _extract_docx(data)
    if extension == ".pptx":
        return _extract_pptx(data)
    if extension == ".xlsx":
        return _extract_xlsx(data)
    if extension == ".rtf":
        return _extract_rtf(data)
    return _decode_text(data)


def _already_indexed(message_id: str, attachment_ref: str) -> bool:
    try:
        conn = sqlite3.connect(DB_PATH, timeout=5.0)
        row = conn.execute(
            "SELECT 1 FROM external_ids WHERE namespace='gmail_attachment' AND external_id=? LIMIT 1",
            (f"{message_id}:{attachment_ref}",),
        ).fetchone()
        conn.close()
        return row is not None
    except Exception:
        return False


def _effective_query(logical_query: str) -> tuple[str, str, int]:
    """Return effective Gmail query, saved page token, and in-progress max internalDate.

    A scan is resumed from Gmail's page token until its result set is exhausted. Once
    complete, future scans overlap the previous high-water mark by 48 hours so equal
    timestamps, delayed delivery, and small clock inconsistencies cannot create a gap.
    """
    pending_token = _state_get("pending_page_token")
    pending_query = _state_get("pending_effective_query")
    pending_logical = _state_get("pending_logical_query")
    pending_max = _state_int("pending_max_internal_ms")
    if pending_token and pending_query and pending_logical == logical_query:
        return pending_query, pending_token, pending_max

    if logical_query != _DEFAULT_QUERY:
        return logical_query, "", 0

    checkpoint_ms = _state_int("last_complete_internal_ms")
    if checkpoint_ms <= 0:
        return logical_query, "", 0
    after_seconds = max(0, checkpoint_ms // 1000 - _OVERLAP_SECONDS)
    return f"{logical_query} after:{after_seconds}", "", 0


def sync(*, query: str = _DEFAULT_QUERY, limit: int = _MAX_MESSAGES) -> dict[str, Any]:
    """Extract passive text from Gmail attachments into the private world graph.

    The scanner is incremental and resumable. It paginates a bounded number of Gmail
    pages per pass, persists a page token when backlog remains, and advances its source
    time checkpoint only when the entire query result has been traversed. A 48-hour
    overlap is intentionally re-read on later passes; stable attachment external IDs
    make that overlap idempotent.

    No attachment is executed or opened through an OS application. Only an explicit
    allowlist of passive text/document formats is parsed in memory. Macro-enabled,
    executable, archive, image and unknown attachment types remain unsupported.
    Image-only/textless PDFs are explicitly reported as needing OCR rather than being
    silently treated as successfully indexed.
    """
    try:
        from jarvis_mrb.tools.google import _load_credentials
        creds = _load_credentials(interactive=False)
    except Exception as exc:
        _state_set(last_status="error", last_error=str(exc)[:500], last_scan_at=datetime.now().astimezone().isoformat())
        try:
            from jarvis_mrb.runtime_health import record_failure
            record_failure("gmail_attachments", exc)
        except Exception:
            pass
        return {"ok": False, "messages": 0, "indexed": 0, "error": str(exc)[:500]}
    if not creds:
        message = "Google is not connected."
        _state_set(last_status="unconfigured", last_error=message, last_scan_at=datetime.now().astimezone().isoformat())
        return {"ok": False, "messages": 0, "indexed": 0, "error": message}

    safe_limit = max(1, min(int(limit), 50))
    logical_query = str(query or _DEFAULT_QUERY).strip() or _DEFAULT_QUERY
    effective_query, page_token, scan_max_internal_ms = _effective_query(logical_query)

    try:
        service = build("gmail", "v1", credentials=creds, cache_discovery=False)
    except Exception as exc:
        _state_set(last_status="error", last_error=str(exc)[:500], last_scan_at=datetime.now().astimezone().isoformat())
        try:
            from jarvis_mrb.runtime_health import record_failure
            record_failure("gmail_attachments", exc)
        except Exception:
            pass
        return {"ok": False, "messages": 0, "indexed": 0, "error": f"Gmail attachment service failed: {exc}"[:500]}

    messages = 0
    indexed = 0
    skipped_type = 0
    skipped_size = 0
    failed_extract = 0
    needs_ocr = 0
    message_failures = 0
    pages = 0
    next_page_token = page_token
    fatal_error = ""

    while pages < _MAX_PAGES_PER_SYNC:
        try:
            request = service.users().messages().list(
                userId="me",
                q=effective_query,
                maxResults=safe_limit,
                pageToken=next_page_token or None,
            )
            listing = request.execute()
        except Exception as exc:
            fatal_error = f"Gmail attachment listing failed: {exc}"[:500]
            break

        pages += 1
        refs = listing.get("messages", []) if isinstance(listing, dict) else []
        for ref in refs:
            message_id = str(ref.get("id") or "").strip() if isinstance(ref, dict) else ""
            if not message_id:
                continue
            try:
                message = service.users().messages().get(userId="me", id=message_id, format="full").execute()
            except Exception:
                message_failures += 1
                continue
            messages += 1
            try:
                internal_ms = int(message.get("internalDate") or 0)
            except (TypeError, ValueError):
                internal_ms = 0
            scan_max_internal_ms = max(scan_max_internal_ms, internal_ms)

            thread_id = str(message.get("threadId") or "").strip()[:500]
            payload = message.get("payload") or {}
            headers = _headers(payload)
            raw_sender = headers.get("from", "")
            display_name, address = parseaddr(raw_sender)
            person_name = " ".join((display_name or address or raw_sender or "Unknown sender").split())[:300]
            subject = " ".join(headers.get("subject", "(no subject)").split())[:500]
            date_raw = str(headers.get("date") or "")[:200]

            sender_id: str | None = None
            if person_name and person_name.lower() != "unknown sender":
                sender_id = ensure_entity(
                    "person",
                    person_name,
                    external_namespace="email_address" if address else None,
                    external_id=address.lower() if address else None,
                    aliases=[address] if address else [],
                    attributes={"email": address.lower()} if address else {},
                    confidence=1.0,
                )
            parent_email_id = ensure_entity(
                "document",
                subject,
                external_namespace="gmail",
                external_id=message_id,
                attributes={"source_kind": "gmail", "sender": raw_sender, "thread_id": thread_id},
                confidence=1.0,
            )

            for part in _iter_parts(payload):
                filename = " ".join(str(part.get("filename") or "").split())[:500]
                if not filename:
                    continue
                extension = PurePath(filename).suffix.lower()
                if extension not in _ALLOWED_EXTENSIONS:
                    skipped_type += 1
                    continue
                body = part.get("body") or {}
                attachment_id = str(body.get("attachmentId") or "").strip()
                part_id = str(part.get("partId") or filename).strip()[:300]
                attachment_ref = attachment_id or part_id
                if _already_indexed(message_id, attachment_ref):
                    continue
                declared_size = int(body.get("size") or 0)
                if declared_size > _MAX_ATTACHMENT_BYTES:
                    skipped_size += 1
                    continue

                try:
                    if attachment_id:
                        response = service.users().messages().attachments().get(
                            userId="me", messageId=message_id, id=attachment_id
                        ).execute()
                        encoded = str(response.get("data") or "")
                    else:
                        encoded = str(body.get("data") or "")
                    data = _decode_base64url(encoded) if encoded else b""
                except Exception:
                    failed_extract += 1
                    continue
                if not data or len(data) > _MAX_ATTACHMENT_BYTES:
                    skipped_size += 1
                    continue

                text = _extract(filename, str(part.get("mimeType") or ""), data)
                if not text:
                    if extension == ".pdf":
                        needs_ocr += 1
                    else:
                        failed_extract += 1
                    continue

                participants: list[tuple[str, str, float]] = [(parent_email_id, "parent_email", 1.0)]
                if sender_id:
                    participants.append((sender_id, "sender", 1.0))
                record_knowledge_source(
                    "gmail_attachment",
                    f"{message_id}:{attachment_ref}",
                    title=filename,
                    text=text,
                    occurred_at=date_raw or None,
                    metadata={
                        "message_id": message_id,
                        "thread_id": thread_id,
                        "attachment_ref": attachment_ref,
                        "filename": filename,
                        "mime_type": str(part.get("mimeType") or "")[:200],
                        "email_subject": subject,
                        "sender": raw_sender,
                        "bytes": len(data),
                        "text_chars": len(text),
                        "gmail_internal_ms": internal_ms,
                    },
                    participants=participants,
                )
                indexed += 1

        next_page_token = str(listing.get("nextPageToken") or "") if isinstance(listing, dict) else ""
        if not next_page_token:
            break

    now = datetime.now().astimezone().isoformat()
    if fatal_error:
        # A page token may expire. Clear it so the next pass restarts from the last
        # fully completed checkpoint rather than getting permanently stuck.
        _state_set(
            pending_page_token="",
            pending_effective_query="",
            pending_logical_query="",
            pending_max_internal_ms=0,
            last_status="error",
            last_error=fatal_error,
            last_scan_at=now,
        )
        try:
            from jarvis_mrb.runtime_health import record_failure
            record_failure("gmail_attachments", RuntimeError(fatal_error))
        except Exception:
            pass
        return {
            "ok": False,
            "messages": messages,
            "indexed": indexed,
            "pages": pages,
            "error": fatal_error,
            "message_failures": message_failures,
        }

    backlog_remaining = bool(next_page_token)
    if backlog_remaining:
        _state_set(
            pending_page_token=next_page_token,
            pending_effective_query=effective_query,
            pending_logical_query=logical_query,
            pending_max_internal_ms=scan_max_internal_ms,
            backlog_remaining=True,
            last_status="backlog",
            last_error="",
            last_scan_at=now,
        )
    else:
        previous_checkpoint = _state_int("last_complete_internal_ms")
        checkpoint = max(previous_checkpoint, scan_max_internal_ms)
        _state_set(
            pending_page_token="",
            pending_effective_query="",
            pending_logical_query="",
            pending_max_internal_ms=0,
            backlog_remaining=False,
            last_complete_internal_ms=checkpoint,
            last_complete_at=now,
            last_status="ok" if not message_failures else "partial",
            last_error="" if not message_failures else f"{message_failures} message fetch failure(s)",
            last_scan_at=now,
        )

    try:
        from jarvis_mrb.runtime_health import record_failure, record_success
        if message_failures:
            record_failure("gmail_attachments", RuntimeError(f"{message_failures} message fetch failure(s)"))
        else:
            record_success("gmail_attachments")
    except Exception:
        pass

    return {
        "ok": True,
        "messages": messages,
        "indexed": indexed,
        "pages": pages,
        "backlog_remaining": backlog_remaining,
        "skipped_unsupported_type": skipped_type,
        "skipped_oversize": skipped_size,
        "failed_or_no_text": failed_extract,
        "unsupported_needs_ocr": needs_ocr,
        "message_failures": message_failures,
        "max_attachment_bytes": _MAX_ATTACHMENT_BYTES,
        "allowed_extensions": sorted(_ALLOWED_EXTENSIONS),
        "checkpoint_internal_ms": _state_int("last_complete_internal_ms"),
        "overlap_seconds": _OVERLAP_SECONDS,
    }


def status() -> dict[str, Any]:
    return {
        "enabled": True,
        "read_only": True,
        "executes_attachments": False,
        "allowed_extensions": sorted(_ALLOWED_EXTENSIONS),
        "max_attachment_bytes": _MAX_ATTACHMENT_BYTES,
        "max_text_chars": _MAX_TEXT_CHARS,
        "preserves_thread_identity": True,
        "resumable_pagination": True,
        "max_pages_per_sync": _MAX_PAGES_PER_SYNC,
        "checkpoint_overlap_seconds": _OVERLAP_SECONDS,
        "last_complete_internal_ms": _state_int("last_complete_internal_ms"),
        "last_complete_at": _state_get("last_complete_at"),
        "last_scan_at": _state_get("last_scan_at"),
        "last_status": _state_get("last_status", "never"),
        "last_error": _state_get("last_error"),
        "backlog_remaining": _state_get("backlog_remaining", "false").lower() in {"true", "1"},
        "pending_page_token": bool(_state_get("pending_page_token")),
        "xlsx_passive_text": True,
        "scanned_pdf_ocr": "explicitly unsupported_needs_ocr",
    }
